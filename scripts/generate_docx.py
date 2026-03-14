#!/usr/bin/env python3
"""Convert a Markdown file to a styled Word (.docx) document.

Usage:
    python generate_docx.py <markdown_file> <output_docx_path>

Supports:
    - Headings (# to ####)
    - Tables (| col1 | col2 |)
    - Bold (**text**), italic (*text*)
    - Bulleted lists (- item / * item)
    - Numbered lists (1. item)
    - Blockquotes (> text)
    - Horizontal rules (---)
    - Page breaks at top-level headings starting with "Part"
"""

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError:
    print("Error: python-docx not installed. Run: pip install python-docx", file=sys.stderr)
    sys.exit(1)


HEADING_COLORS = {
    0: RGBColor(0x1A, 0x5C, 0x8E),  # Title
    1: RGBColor(0x2C, 0x3E, 0x50),  # H1
    2: RGBColor(0x34, 0x49, 0x5E),  # H2
    3: RGBColor(0x5D, 0x6D, 0x7E),  # H3
    4: RGBColor(0x7F, 0x8C, 0x8D),  # H4
}


def add_heading(doc, text, level):
    """Add a colored heading."""
    heading = doc.add_heading(text, level=level)
    color = HEADING_COLORS.get(level, HEADING_COLORS[4])
    for run in heading.runs:
        run.font.color.rgb = color
    return heading


def add_rich_paragraph(doc, text, style=None):
    """Add a paragraph with inline bold/italic formatting."""
    para = doc.add_paragraph(style=style)
    # Parse inline formatting: **bold**, *italic*, `code`
    pattern = r'(\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`)'
    last_end = 0
    for match in re.finditer(pattern, text):
        # Add text before this match
        if match.start() > last_end:
            para.add_run(text[last_end:match.start()])
        if match.group(2):  # bold
            run = para.add_run(match.group(2))
            run.bold = True
        elif match.group(3):  # italic
            run = para.add_run(match.group(3))
            run.italic = True
        elif match.group(4):  # code
            run = para.add_run(match.group(4))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        last_end = match.end()
    # Add remaining text
    if last_end < len(text):
        para.add_run(text[last_end:])
    return para


def parse_table(lines, start_idx):
    """Parse a markdown table starting at start_idx. Returns (rows_data, end_idx)."""
    rows = []
    i = start_idx
    while i < len(lines):
        line = lines[i].strip()
        if not line.startswith("|"):
            break
        # Skip separator row (|---|---|)
        if re.match(r'^\|[\s\-:]+\|', line):
            i += 1
            continue
        cells = [c.strip() for c in line.split("|")[1:-1]]
        rows.append(cells)
        i += 1
    return rows, i


def add_table(doc, rows):
    """Add a table to the document."""
    if not rows or not rows[0]:
        return
    n_cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=n_cols)
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            if j >= n_cols:
                break
            cell = t.cell(i, j)
            cell.text = ""
            # Strip markdown bold from cell text
            clean = re.sub(r'\*\*(.+?)\*\*', r'\1', cell_text)
            run = cell.paragraphs[0].add_run(clean)
            run.font.size = Pt(10)
            run.font.name = "Calibri"
            if i == 0:  # header row
                run.bold = True
    doc.add_paragraph()


def md_to_docx(md_path, output_path):
    """Convert a Markdown file to a Word document."""
    md_text = Path(md_path).read_text(encoding="utf-8")
    lines = md_text.split("\n")

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Empty line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            # Add a thin horizontal line
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(6)
            pPr = para._p.get_or_add_pPr()
            pBdr = pPr.makeelement(qn('w:pBdr'), {})
            bottom = pBdr.makeelement(qn('w:bottom'), {
                qn('w:val'): 'single',
                qn('w:sz'): '6',
                qn('w:space'): '1',
                qn('w:color'): 'CCCCCC',
            })
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,4})\s+(.+)', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            # Page break before Part headings (level 0 equivalent = # Part ...)
            if level == 1 and text.lower().startswith("part"):
                doc.add_page_break()
                add_heading(doc, text, level=0)
            else:
                add_heading(doc, text, level=level)
            i += 1
            continue

        # Table
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue

        # Blockquote
        if stripped.startswith(">"):
            quote_text = stripped.lstrip("> ").strip()
            para = add_rich_paragraph(doc, quote_text)
            pf = para.paragraph_format
            pf.left_indent = Inches(0.5)
            for run in para.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
            i += 1
            continue

        # Bulleted list
        bullet_match = re.match(r'^[\-\*]\s+(.+)', stripped)
        if bullet_match:
            add_rich_paragraph(doc, bullet_match.group(1), style="List Bullet")
            i += 1
            continue

        # Numbered list
        num_match = re.match(r'^\d+\.\s+(.+)', stripped)
        if num_match:
            add_rich_paragraph(doc, num_match.group(1), style="List Number")
            i += 1
            continue

        # Regular paragraph
        add_rich_paragraph(doc, stripped)
        i += 1

    doc.save(str(output_path))
    return str(output_path)


def main():
    if len(sys.argv) < 3:
        print("Usage: generate_docx.py <markdown_file> <output_docx_path>", file=sys.stderr)
        sys.exit(1)

    md_path = sys.argv[1]
    output_path = sys.argv[2]

    if not Path(md_path).exists():
        print(f"Error: File not found: {md_path}", file=sys.stderr)
        sys.exit(1)

    result = md_to_docx(md_path, output_path)
    print(f"Word document saved: {result}")


if __name__ == "__main__":
    main()
